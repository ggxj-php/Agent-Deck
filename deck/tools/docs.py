"""文档工具：读工作区里的 PDF / Word / 图片。

PDF 和 Word 里的**文字**直接抽出来；抽不出文字的图（扫描页、插图、照片）可以
顺手丢给视觉模型看一眼，返回图片描述——走的是和聊天里直接发图**同一条多模态链路**
（`settings.vision_model()` + `key_for(vision=True)` + `base_url_for(vision=True)`），
所以后台配的图片模型 / key / 接口地址一变，这里也跟着变。

看图按张计费，所以默认每次最多描述 3 张；额度用完只报还剩几张，不再调用。
"""
import base64
import os

from .. import llm, settings
from ..sandbox import safe_path

MAX_OUTPUT_CHARS = 20000          # 一次返回给模型的文字上限
MAX_IMAGE_BYTES = 4 * 1024 * 1024  # 和聊天发图的上限保持一致
MIN_IMAGE_BYTES = 6 * 1024         # 太小的多半是图标 / 装饰，不值得花一次调用
DEFAULT_MAX_IMAGES = 3             # 单次最多描述几张图
HARD_MAX_IMAGES = 10
DEFAULT_PAGE_CAP = 12              # 不给 pages 时最多读多少页
VISION_MAX_TOKENS = 600

DEFAULT_IMAGE_PROMPT = (
    "用中文简要描述这张图的内容。如果有文字，把关键文字也抄出来；"
    "如果是图表，说明它画的是什么、有哪些关键数据。200 字以内，不要客套。"
)


def _sniff_mime(data):
    """靠魔数认图片格式，别信扩展名。"""
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith(b"GIF87a") or data.startswith(b"GIF89a"):
        return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data.startswith(b"BM"):
        return "image/bmp"
    return ""


def _prepare_image(data):
    """压缩一下再送模型：长边超 1600 就缩小，统一转 JPEG。

    文档里抽出来的图经常是几 MB 的位图，直接发又慢又容易撞上 4MB 上限；
    缩到 1600 长边对「看图说话」完全够用，顺便省 token。
    没装 Pillow 或者压不了就原样返回，不影响主流程。
    """
    try:
        import io

        from PIL import Image

        img = Image.open(io.BytesIO(data))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        if max(img.size) > 1600:
            img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=85)
        out = buf.getvalue()
        return out if len(out) < len(data) else data
    except Exception:                              # noqa: BLE001
        return data


def _describe_image(data, prompt=""):
    """把一张图丢给视觉模型，返回 (描述, 错误)。"""
    if not data:
        return None, "图片数据是空的"
    data = _prepare_image(data)
    mime = _sniff_mime(data)
    if not mime:
        return None, "格式认不出来（不是常见图片）"
    if len(data) > MAX_IMAGE_BYTES:
        return None, f"{len(data)} 字节，超过 {MAX_IMAGE_BYTES} 上限（和聊天发图一致）"
    url = "data:" + mime + ";base64," + base64.b64encode(data).decode("ascii")
    ask = (prompt or "").strip() or DEFAULT_IMAGE_PROMPT
    try:
        text = llm.complete(
            [{
                "role": "user",
                "content": [
                    {"type": "text", "text": ask},
                    {"type": "image_url", "image_url": {"url": url}},
                ],
            }],
            model=settings.vision_model(),
            api_key=settings.key_for(vision=True),
            base_url=settings.base_url_for(vision=True),
            max_tokens=VISION_MAX_TOKENS,
            thinking=False,
        )
    except Exception as exc:                       # noqa: BLE001
        return None, f"视觉模型调用失败：{exc}"
    text = (text or "").strip()
    if not text:
        return None, "视觉模型没返回内容"
    return text, None


def _parse_pages(spec, total):
    """解析 pages 参数：""、"3"、"1-5"、"1,4,7"、"2-4,9"。"""
    if not total:
        return []
    spec = (spec or "").strip()
    if not spec:
        return list(range(1, min(total, DEFAULT_PAGE_CAP) + 1))
    picked = set()
    for piece in spec.replace("，", ",").split(","):
        piece = piece.strip()
        if not piece:
            continue
        if "-" in piece:
            lo, _, hi = piece.partition("-")
            try:
                lo = int(lo.strip())
                hi = int(hi.strip())
            except ValueError:
                continue
            for n in range(max(1, lo), min(total, hi) + 1):
                picked.add(n)
        else:
            try:
                n = int(piece)
            except ValueError:
                continue
            if 1 <= n <= total:
                picked.add(n)
    return sorted(picked)


def _cap(text, budget):
    """总字数封顶，超了截断并说明。"""
    if len(text) <= budget:
        return text, budget - len(text)
    keep = text[:budget]
    return keep + f"\n…（内容太长，后面省略了；总长 {len(text)} 字）", 0


def read_img(path: str, prompt: str = "") -> str:
    """读工作区里的一张图片，让视觉模型描述它。"""
    full = safe_path(path)
    if not os.path.isfile(full):
        return f"错误：文件不存在 {path}"
    size = os.path.getsize(full)
    if size > MAX_IMAGE_BYTES:
        return f"错误：图片 {size} 字节，超过 {MAX_IMAGE_BYTES} 上限（和聊天发图一致，先压小一点）"
    if size < MIN_IMAGE_BYTES:
        return f"错误：图片只有 {size} 字节，太小了，多半不是正经图"
    try:
        with open(full, "rb") as fh:
            data = fh.read()
    except OSError as exc:
        return f"读不了这个文件：{exc}"
    desc, err = _describe_image(data, prompt)
    if err:
        return f"读图失败：{err}"
    return f"【{path}】\n{desc}"


def read_pdf(path: str, pages: str = "", with_images: bool = True,
             max_images: int = DEFAULT_MAX_IMAGES) -> str:
    """读工作区里的 PDF：按页抽文字，抽不出字的图可以交给视觉模型看。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "错误：这台机器没装 pypdf，读不了 PDF"

    full = safe_path(path)
    if not os.path.isfile(full):
        return f"错误：文件不存在 {path}"
    try:
        reader = PdfReader(full)
        total = len(reader.pages)
    except Exception as exc:                       # noqa: BLE001
        return f"打不开这个 PDF（{path}）：{exc}"
    if not total:
        return f"这个 PDF 一页都没有：{path}"

    wanted = _parse_pages(pages, total)
    if not wanted:
        return f"pages 参数没解析出有效页码（这个 PDF 共 {total} 页）"

    budget = MAX_OUTPUT_CHARS
    img_budget = max(0, min(int(max_images or 0), HARD_MAX_IMAGES)) if with_images else 0
    skipped_images = 0
    lines = [f"《{path}》共 {total} 页，本次读第 {wanted[0]}~{wanted[-1]} 页（{len(wanted)} 页）"
             + ("，图会顺手让视觉模型看一眼" if img_budget else "，跳过图")]

    for number in wanted:
        page = reader.pages[number - 1]
        try:
            text = (page.extract_text() or "").strip()
        except Exception:                          # noqa: BLE001
            text = ""
        images = []
        try:
            images = list(page.images)
        except Exception:                          # noqa: BLE001
            images = []

        lines.append(f"\n--- 第 {number} 页 ---")
        if text:
            lines.append(text)
        elif images:
            lines.append(f"（这一页没有文字层，是图片 / 扫描页，共 {len(images)} 张图）")
        else:
            lines.append("（这一页既没有文字也没有图）")

        if not images:
            continue
        for idx, img in enumerate(images, 1):
            try:
                data = img.data
            except Exception:                      # noqa: BLE001
                continue
            if len(data) < MIN_IMAGE_BYTES:
                skipped_images += 1
                continue
            if img_budget <= 0:
                skipped_images += 1
                continue
            img_budget -= 1
            desc, err = _describe_image(data)
            lines.append(f"[图 {idx}] " + (desc if desc else f"（没看成：{err}）"))

    if skipped_images:
        lines.append(f"\n（另有 {skipped_images} 张图没看：额度用完或图片太小；"
                     f"要看得把 max_images 调大一点）")
    return _cap("\n".join(lines), budget)[0]


def read_doc(path: str, with_images: bool = True,
             max_images: int = DEFAULT_MAX_IMAGES) -> str:
    """读工作区里的 Word（.docx）：段落 + 表格 + 内嵌图。"""
    try:
        import docx
    except ImportError:
        return "错误：这台机器没装 python-docx，读不了 Word"

    full = safe_path(path)
    if not os.path.isfile(full):
        return f"错误：文件不存在 {path}"
    try:
        document = docx.Document(full)
    except Exception as exc:                       # noqa: BLE001
        return f"打不开这个 Word（{path}）：{exc}（老式 .doc 不支持，先另存成 .docx）"

    parts = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join((cell.text or "").strip() for cell in row.cells))
    body = "\n".join(parts).strip()

    blobs = []
    try:
        for related in document.part.related_parts.values():
            ctype = getattr(related, "content_type", "") or ""
            if not ctype.startswith("image/"):
                continue
            try:
                blobs.append(related.blob)
            except Exception:                      # noqa: BLE001
                continue
    except Exception:                              # noqa: BLE001
        blobs = []

    lines = [f"《{path}》：正文 {len(body)} 字，内嵌图 {len(blobs)} 张"]
    lines.append("\n--- 正文 ---")
    lines.append(body or "（没有文字内容）")

    img_budget = max(0, min(int(max_images or 0), HARD_MAX_IMAGES)) if with_images else 0
    skipped = 0
    if blobs:
        lines.append("\n--- 内嵌图 ---")
        for idx, data in enumerate(blobs, 1):
            if len(data) < MIN_IMAGE_BYTES or img_budget <= 0:
                skipped += 1
                continue
            img_budget -= 1
            desc, err = _describe_image(data)
            lines.append(f"[图 {idx}] " + (desc if desc else f"（没看成：{err}）"))
        if skipped:
            lines.append(f"（另有 {skipped} 张图没看：额度用完或图片太小）")
    return _cap("\n".join(lines), MAX_OUTPUT_CHARS)[0]


SPECS = [
    {
        "type": "function",
        "function": {
            "name": "read_pdf",
            "description": (
                "读取工作区 workspace 里的 PDF，按页抽取文字（带页码）。"
                "扫描页 / 插图这类抽不出文字的图，会交给视觉模型看一眼并返回描述。"
                "看图按张计费，默认最多 3 张，可以用 max_images 调整。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "相对 workspace 的文件路径，例如 books/foo.pdf",
                    },
                    "pages": {
                        "type": "string",
                        "description": (
                            "要读哪些页，支持 \"3\"、\"1-5\"、\"1,4,7\"；"
                            "不填就默认读前 12 页"
                        ),
                    },
                    "with_images": {
                        "type": "boolean",
                        "description": "是否把图交给视觉模型看，默认 true；只想要文字就传 false",
                    },
                    "max_images": {
                        "type": "integer",
                        "description": "本次最多看几张图，默认 3，上限 10",
                    },
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_doc",
            "description": (
                "读取工作区 workspace 里的 Word 文档（.docx），抽取段落、表格和内嵌图片。"
                "内嵌图会交给视觉模型看一眼并返回描述，默认最多 3 张。"
                "老式 .doc 不支持，得先另存成 .docx。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "相对 workspace 的文件路径，例如 reports/周报.docx",
                    },
                    "with_images": {
                        "type": "boolean",
                        "description": "是否把内嵌图交给视觉模型看，默认 true",
                    },
                    "max_images": {
                        "type": "integer",
                        "description": "本次最多看几张图，默认 3，上限 10",
                    },
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_img",
            "description": (
                "读取工作区 workspace 里的一张图片，交给视觉模型描述内容"
                "（图里的文字、图表含义、画面内容）。和聊天里直接发图走同一套多模态配置。"
                "支持 png/jpg/gif/webp/bmp，单张不超过 4MB。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "相对 workspace 的图片路径，例如 samples/chart.png",
                    },
                    "prompt": {
                        "type": "string",
                        "description": "想重点看什么（可选），例如「这张表里第三列是什么」",
                    },
                },
                "required": ["path"],
            },
        },
    },
]
